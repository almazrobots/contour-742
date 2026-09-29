"""Протокол по образцу Приложения № 2 (OS-INSP-5.1.6, T-121): разделы в порядке образца, экранирование разметки (OWASP H6)."""

import io

import pypdfium2 as pdfium
import pytest
from docx import Document

from inspector_ml.render_appendix2 import render_docx, render_pdf

pytestmark = [pytest.mark.l1_functional, pytest.mark.l6_adversarial]

TITLES = [
    "Раздел 1. Статус загрузки документов",
    "Раздел 2. Сводная статистика",
    "Раздел 3. Параметры, не проверенные из-за отсутствия ИД — 1",
    "Раздел 4. Критические нарушения — 1",
    "Раздел 5. Существенные нарушения — 0",
    "Раздел 6. Подозрения ИИ — 1",
    "Раздел 7. Резолютивная часть",
]
V = ["№", "Раздел", "Параметр (код)", "ПД", "РД", "ИД", "Отклонение", "Решение инспектора"]


def sample(description: str = "Класс указан по-разному") -> dict:
    return {
        "protocol_version": 6,
        "appendix2": {
            "title": "Протокол автоматизированной сверки № P-20260927-6fa6c74a",
            "header": [["Объект", "Алтуфьевское ш., 79Б"], ["Адрес", "—"], ["Номер надзорного дела", "—"], ["Застройщик", "—"], ["Подрядчик", "—"],
                       ["Дата формирования", "27 сентября 2026 г."], ["Версия протокола", "6 (предварительная)"], ["Статус", "Верификация завершена"]],
            "sections": [
                {"title": TITLES[0], "columns": ["Тип документа", "Статус", "Загружено файлов", "Ожидается", "Комментарий"], "widths": [16, 18, 16, 14, 36],
                 "rows": [["ПД", "Частично", "36", "38", "Отсутствует файлов: 2"], ["РД", "Полностью", "21", "21", "Все файлы загружены"], ["ИД", "Не требуется", "0", "0", "Не требуются по реестру"]]},
                {"title": TITLES[1], "columns": ["Показатель", "Количество", "% от общего"], "widths": [60, 20, 20], "rows": [["Всего параметров в проверке", "132", "100%"]]},
                {"title": TITLES[2], "columns": ["№", "Код", "Раздел", "Параметр", "Нужный документ ИД (по Матрице)"], "widths": [5, 12, 10, 35, 38],
                 "rows": [["1", "KR-055", "КР", "Класс прочности бетона", "Протоколы испытаний образцов"]]},
                {"title": TITLES[3], "columns": V, "widths": [5, 8, 30, 12, 12, 12, 10, 11], "rows": [["1", "КР", "Класс бетона (KR-055)", "B35", "B25", "—", "—", "Ожидает"]]},
                {"title": TITLES[4], "columns": V, "widths": [5, 8, 30, 12, 12, 12, 10, 11], "rows": []},
                {"title": TITLES[5], "columns": ["№", "Метод", "Описание", "ПД", "РД", "ИД", "Решение инспектора", "Комментарий"], "widths": [4, 10, 27, 25, 12, 5, 8, 9],
                 "rows": [["1", "Внутреннее противоречие", description, "ПБ стр. 7", "—", "—", "Уточнение", "—"]]},
                {"title": TITLES[6], "columns": [], "widths": [], "rows": []},
            ],
            "resolution_columns": ["№", "Нарушение (код)", "Рекомендация инспектора", "Основание по Матрице"],
            "resolution": [{"title": "7.1. По критическим нарушениям", "rows": [["1", "Класс бетона (KR-055)", "—", "Понижение класса бетона"]]},
                           {"title": "7.2. По существенным нарушениям", "rows": []}],
            "note": "Уровень риска определяет только очерёдность экспертной проверки.",
        },
    }


def pdf_text(b: bytes) -> str:
    doc = pdfium.PdfDocument(b)
    return "\n".join(doc[i].get_textpage().get_text_range() for i in range(len(doc)))


def test_pdf_sections_in_appendix2_order():
    text = pdf_text(render_pdf(sample()))
    pos = [text.find(t.split(". ", 1)[1].split(" — ")[0]) for t in TITLES]
    assert all(p >= 0 for p in pos), pos
    assert pos == sorted(pos)
    assert "KR-055" in text and "Нет записей" in text and "стр. 1" in text


def test_pdf_escapes_markup_from_data():
    # OWASP H6: «<link>» и «<b>» в описании — текст, а не разметка; рендер не падает
    text = pdf_text(render_pdf(sample('Проём <b>2,5 м</b> & <link href="http://evil">ссылка</link>')))
    assert "<b>2,5 м</b>" in text.replace("\r", "").replace("\n", " ") or "<b>2,5" in text
    assert "evil" in text  # напечатано как текст


def test_docx_same_sections_and_rows():
    d = Document(io.BytesIO(render_docx(sample())))
    paras = "\n".join(p.text for p in d.paragraphs)
    for t in TITLES:
        assert t in paras
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "Класс бетона (KR-055)" in cells and "Уточнение" in cells


def test_docx_survives_control_characters_from_documents():
    # OWASP H6: управляющие символы из текста документа недопустимы в XML 1.0 — python-docx падал бы на сохранении
    d = Document(io.BytesIO(render_docx(sample("Класс\x00 С1\x0b\x1f по ПБ"))))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "Класс С1 по ПБ" in cells
