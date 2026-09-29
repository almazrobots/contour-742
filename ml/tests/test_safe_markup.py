"""Экранирование текста протокола PDF/DOCX (OS-INSP-5.1.20, 5.1.21; T-094, OWASP H6).

Враждебные значения из документов, ответов ИИ и комментариев должны попадать в протокол как текст:
разметка ReportLab не исполняется, локальные файлы не встраиваются, недопустимые в XML символы не роняют DOCX.
"""

import io

import pytest

from inspector_ml.safe_markup import docx_text, pdf_text


def _protocol(value: str) -> dict:
    """Синтетический протокол: враждебное значение во всех полях из данных — шапке и каждом разделе."""
    src = {
        "stage": "PD",
        "document_code": value,
        "revision": "1",
        "page": 1,
        "bbox": None,
    }
    card = {
        "finding_id": "F-1",
        "parameter_name": value,
        "expected": value,
        "actual": value,
        "delta": value,
        "sources": [src],
        "decision": {"user_name": value, "comment": value},
    }
    return {
        "process_id": "P-1",
        "protocol_version": 1,
        "status": "READY",
        "object": {"name": value, "address": value, "permit_number": value},
        "check_type": {"scenario": "S", "title": value},
        "upload_status": {value: "x"},
        "versions": {
            "matrix_version": "1",
            "model_version": "m",
            "dataset_version": "d",
            "input_manifest_hash": value,
        },
        "sections": {
            "completeness": [
                {
                    "param_code": "P1",
                    "parameter_name": value,
                    "status": value,
                    "reason": value,
                }
            ],
            "candidates": [card],
            "confirmed_violations": [card],
            "negative_verified": [
                {
                    "param_code": "P2",
                    "parameter_name": value,
                    "expected": value,
                    "actual": value,
                }
            ],
            "suspicions": [
                {
                    "discovery_method": value,
                    "description": value,
                    "normative_base": value,
                    "review_priority": value,
                }
            ],
        },
    }


def _pdf_text_and_images(pdf: bytes) -> tuple[str, int]:
    """Текст всех страниц (без переносов строк) и число встроенных картинок (XObject Image)."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    doc = pdfium.PdfDocument(pdf)
    text, images = [], 0
    for page in doc:
        text.append(page.get_textpage().get_text_range())
        images += sum(
            1 for _ in page.get_objects(filter=[raw.FPDF_PAGEOBJ_IMAGE], max_depth=5)
        )
    return "".join(text).replace("\r", "").replace("\n", ""), images


# --- сами функции -------------------------------------------------------------------------------


@pytest.mark.l1_functional
def test_pdf_text_экранирует_служебные_символы_reportlab():
    """OS-INSP-5.1.20: < > & превращаются в сущности, которые Paragraph покажет как написаны."""
    assert pdf_text("a < b & c > d") == "a &lt; b &amp; c &gt; d"
    assert (
        pdf_text('<font color="white">9</font>')
        == '&lt;font color="white"&gt;9&lt;/font&gt;'
    )


@pytest.mark.l3_boundary
def test_pdf_text_none_число_пустая_кириллица():
    assert pdf_text(None) == ""
    assert pdf_text(12.5) == "12.5"
    assert pdf_text(0) == "0"
    assert pdf_text("") == ""
    assert pdf_text("Толщина стены, мм — 380") == "Толщина стены, мм — 380"


@pytest.mark.l1_functional
def test_docx_text_убирает_недопустимые_в_xml_символы():
    """OS-INSP-5.1.21: управляющие символы уходят, табуляция и переводы строк остаются."""
    assert docx_text("x\x00y\x0bz\x1f") == "xyz"
    assert docx_text("a\tb\nc\rd") == "a\tb\nc\rd"
    bad = "".join(chr(c) for c in [*range(0x00, 0x09), 0x0B, 0x0C, *range(0x0E, 0x20)])
    assert docx_text("[" + bad + "]") == "[]"
    assert docx_text("a\ud800b\udfffc￾d￿e") == "abcde"
    # граничные допустимые символы соседей по диапазонам сохраняются
    assert docx_text("\x20퟿�\U00010000") == "\x20퟿�\U00010000"


@pytest.mark.l3_boundary
def test_docx_text_none_число_пустая_кириллица():
    assert docx_text(None) == ""
    assert docx_text(7) == "7"
    assert docx_text("") == ""
    assert (
        docx_text("Толщина стены, мм — 380 <b> & >")
        == "Толщина стены, мм — 380 <b> & >"
    )


# --- протокол PDF (OS-INSP-5.1.20) ---------------------------------------------------------------


@pytest.mark.l6_adversarial
def test_pdf_протокола_со_знаками_меньше_больше_амперсанд_выгружается_и_показывает_их_как_есть():
    from inspector_ml.render import render_pdf

    text, _ = _pdf_text_and_images(render_pdf(_protocol("a < b & c > d")))
    assert "a < b & c > d" in text
    assert text.count("a < b & c > d") >= 10  # шапка и все разделы, а не одно поле
    # без экранирования: «a<b» — незакрытый тег (500), «&amp;» и «&#65;» — сущности, «<a href>» — ссылка
    for value in ("a<b", "x &amp; y &#65;", "<a href='file:///etc/hosts'>ссылка</a>"):
        text, _ = _pdf_text_and_images(render_pdf(_protocol(value)))
        assert value in text


@pytest.mark.l6_adversarial
def test_pdf_протокола_тег_font_из_данных_не_исполняется_и_виден_буквально():
    from inspector_ml.render import render_pdf

    value = '<font color="white">999</font>'
    text, _ = _pdf_text_and_images(render_pdf(_protocol(value)))
    assert value in text


@pytest.mark.l6_adversarial
def test_pdf_протокола_тег_img_из_данных_не_встраивает_локальный_файл(tmp_path):
    from PIL import Image

    from inspector_ml.render import render_pdf

    # настоящая картинка на диске: без экранирования ReportLab встроил бы её в PDF
    png = tmp_path / "secret.png"
    Image.new("RGB", (4, 4), "red").save(png)
    for value in (
        '<img src="/etc/hosts"/>',
        f'<img src="{png}" width="20" height="20"/>',
    ):
        text, images = _pdf_text_and_images(render_pdf(_protocol(value)))
        assert value.replace(" ", "") in text.replace(" ", "")
        assert images == 0


@pytest.mark.l6_adversarial
def test_pdf_протокола_служебная_разметка_заголовков_не_ломается():
    """Экранируется только подставляемое значение: наши подписи колонок и заголовки разделов на месте."""
    from inspector_ml.render import render_pdf

    text, _ = _pdf_text_and_images(render_pdf(_protocol("&")))
    assert "1. Комплектность и сопоставимость — 1" in text
    assert "&lt;" not in text and "&amp;" not in text


# --- протокол DOCX (OS-INSP-5.1.21) --------------------------------------------------------------


def _docx_texts(data: bytes) -> list[str]:
    from docx import Document

    d = Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        out += [c.text for row in t.rows for c in row.cells]
    return out


@pytest.mark.l6_adversarial
def test_docx_протокола_с_управляющими_символами_собирается_и_открывается():
    from inspector_ml.render import render_docx

    texts = _docx_texts(render_docx(_protocol("x\x00y\x0bz\x1f")))
    assert sum(t == "xyz" for t in texts) >= 10  # шапка и все разделы
    assert not any(ch in t for t in texts for ch in "\x00\x0b\x1f")


@pytest.mark.l6_adversarial
def test_docx_протокола_сохраняет_табуляцию_перевод_строки_и_знаки_разметки():
    from inspector_ml.render import render_docx

    texts = _docx_texts(render_docx(_protocol("a\tb\nc <i> & d")))
    assert "a\tb\nc <i> & d" in texts


@pytest.mark.l6_adversarial
def test_pdf_протокола_с_длинным_неразрывным_значением_выгружается():
    """Длинный путь/хеш без пробелов не отбирает ширину у соседних колонок и не роняет выгрузку."""
    from inspector_ml.render import render_pdf

    text, _ = _pdf_text_and_images(render_pdf(_protocol("/app/var/blobs/" + "ab" * 200)))
    assert "ab" * 20 in text


def _only_header(value: str) -> dict:
    """Враждебное значение только в шапке, разделы пусты — шапка проверяется сама по себе."""
    p = _protocol(value)
    p["sections"] = {k: [] for k in p["sections"]}
    return p


@pytest.mark.l6_adversarial
def test_pdf_протокола_шапка_показывает_разметку_из_данных_буквально():
    from inspector_ml.render import render_pdf

    value = '<font color="white">999</font> & <b>'
    text, _ = _pdf_text_and_images(render_pdf(_only_header(value)))
    assert "Хеш реестра входных файлов " + value in text
    assert "Объект " + value + ", " + value in text


@pytest.mark.l6_adversarial
def test_docx_протокола_шапка_без_управляющих_символов_и_с_подписями():
    from inspector_ml.render import render_docx

    texts = _docx_texts(render_docx(_only_header("x\x00y\x1f")))
    assert texts[texts.index("Хеш реестра входных файлов") + 1] == "xy"
    assert texts[texts.index("Объект") + 1] == "xy, xy"


@pytest.mark.l1_functional
def test_ширины_колонок_протокола_делят_страницу_узкие_вдвое_уже():
    from inspector_ml.render import _col_widths

    w = _col_widths(["Код", "Параметр", "Статус", "Основание"], 300.0)
    assert w == [50.0, 100.0, 50.0, 100.0]
