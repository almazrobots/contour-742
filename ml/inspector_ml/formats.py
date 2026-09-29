"""Форматы сверх ТЗ §9.1 — XLSX и изображения JPG, PNG, TIF (GAP-INSP-04, OS-INSP-1.2.9).

В реальном комплекте ИД 36 таблиц (ведомости, реестры, АОСР в Excel) и 10 изображений (сканы
паспортов и сертификатов) — см. docs/gera/inspector/corpus-stats.md.

XLSX: страница = лист книги, строка = строка листа, «слова» строки = непустые ячейки. Геометрии нет
(bbox = None), как у DOCX. Изображение: каждая страница (кадр TIFF) уходит в тот же OCR, что
страница PDF без текстового слоя, — через `parse._ocr_page` (ядро OCR не дублируется).
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from pathlib import Path

from PIL import Image, ImageOps, ImageSequence, UnidentifiedImageError

from .model import Line, Page, ParsedDoc, Word
from .normalize import nfc
from .resource_scope import render_guarded

DEFAULT_DPI = (
    300  # у изображения без метаданных dpi считаем, что оно отсканировано в 300 dpi
)

# ─────────────────────────────────────────────── сигнатуры


def image_kind(head: bytes) -> str | None:
    """JPEG, PNG, TIFF по сигнатуре; иначе None."""
    if head.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if head.startswith((b"II*\x00", b"MM\x00*")):
        return "tif"
    return None


def is_xlsx(path: Path) -> bool:
    """ZIP-контейнер OOXML с книгой Excel (xl/workbook.xml), а не документ Word."""
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
    except zipfile.BadZipFile:
        return False
    return any(n.startswith("xl/") for n in names) and not any(
        n.startswith("word/") for n in names
    )


# ─────────────────────────────────────────────── XLSX


def _cell_text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, float):
        # 12450.0 → «12450»; 1.2 → «1,2» (русская запись, как в документах)
        return str(int(v)) if v.is_integer() else repr(v).replace(".", ",")
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y")
    return nfc(str(v))


def parse_xlsx(path: Path, sha256: str) -> ParsedDoc:
    from openpyxl import load_workbook

    try:
        # файловый объект, а не путь: блоб хранится под именем-хешем без расширения, а openpyxl
        # по пути проверяет расширение .xlsx
        wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    except (
        Exception
    ) as e:  # openpyxl: InvalidFileException, BadZipFile, KeyError на битой книге
        raise _corrupted(f"повреждённый XLSX: {path.name}") from e
    pages: list[Page] = []
    try:
        for i, ws in enumerate(wb.worksheets, start=1):
            lines: list[Line] = []
            for row in ws.iter_rows(values_only=True):
                cells = [t for t in (_cell_text(v) for v in row) if t]
                if cells:
                    lines.append(
                        Line(text=" ".join(cells), words=[Word(text=c) for c in cells])
                    )
            pages.append(
                Page(page=i, width=0, height=0, source="structured", lines=lines)
            )
    finally:
        wb.close()
    return ParsedDoc(sha256=sha256, kind="xlsx", engine="openpyxl", pages=pages)


# ─────────────────────────────────────────────── изображения


class _ImagePage:
    """Изображение в роли страницы pdfium для `parse._ocr_page`: размер в пунктах и рендер в заданном масштабе."""

    def __init__(self, img: Image.Image, dpi: float):
        self.img, self.dpi = img, dpi

    def get_size(self) -> tuple[float, float]:
        return self.img.width * 72 / self.dpi, self.img.height * 72 / self.dpi

    def get_rotation(self) -> int:
        return 0  # ориентация EXIF уже применена

    def render(self, scale: float = 1.0, **_: object) -> "_ImagePage":
        w, h = self.get_size()
        size = (max(1, round(w * scale)), max(1, round(h * scale)))
        self._rendered = (
            self.img if size == self.img.size else render_guarded(
                lambda: self.img.resize(size, Image.LANCZOS), *size, 1
            )
        )
        return self

    def to_pil(self) -> Image.Image:
        return self._rendered


def _dpi(img: Image.Image) -> float:
    d = img.info.get("dpi")
    try:
        x = float(d[0]) if d else 0.0
    except (TypeError, ValueError, IndexError):
        x = 0.0
    return x if 50 <= x <= 1200 else DEFAULT_DPI


def parse_image(path: Path, sha256: str) -> ParsedDoc:
    from . import parse  # поздний импорт: parse импортирует этот модуль в detect_kind

    try:
        img = Image.open(str(path))
        try:
            render_guarded(lambda: (img.load(), img)[1], img.width, img.height, 1)
        except BaseException:
            img.close()
            raise
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as e:
        raise _corrupted(f"повреждённое изображение: {path.name}") from e
    pages: list[Page] = []
    try:
        for i, frame in enumerate(ImageSequence.Iterator(img), start=1):
            f = _copy_frame(frame)
            pages.append(parse._ocr_page(_ImagePage(f, _dpi(img)), i))
            del f
    finally:
        img.close()
    engines = {"pillow"} | (
        {"tesseract"} if any(p.quality != "ABSTAIN" for p in pages) else set()
    )
    return ParsedDoc(
        sha256=sha256, kind="image", engine="+".join(sorted(engines)), pages=pages
    )


def _copy_frame(frame: Image.Image) -> Image.Image:
    def copy():
        f = ImageOps.exif_transpose(frame.copy())
        if f.mode not in ("L", "RGB"):
            f = f.convert("RGB" if f.mode in ("RGBA", "P", "CMYK", "LA") else "L")
        return f

    return render_guarded(copy, frame.width, frame.height, 1)


def image_page(path: Path, number: int, dpi: float) -> Image.Image | None:
    """Кадр number (с 1) многостраничного изображения — как его видел OCR при разборе (parse_image)."""
    with Image.open(str(path)) as img:
        for i, frame in enumerate(ImageSequence.Iterator(img), start=1):
            if i == number:
                f = _copy_frame(frame)
                return _ImagePage(f, _dpi(img)).render(scale=dpi / 72).to_pil()
    return None


def _corrupted(msg: str) -> Exception:
    from .parse import CorruptedFile

    return CorruptedFile(msg)
