"""Разбор документа в страницы → строки → слова с координатами (OS-INSP-2.1).

PDF: текстовый слой через pdfium; страница без текстового слоя уходит в OCR (OS-INSP-2.1.1).
Координаты нормируются через FPDF_PageToDevice — он учитывает CropBox и /Rotate (OS-INSP-2.2.2).
DOCX и XML — структурированный текст без геометрии (bbox = None, page = 1).
"""

from __future__ import annotations

import os
from concurrent.futures import ThreadPoolExecutor

import ctypes
import threading
import defusedxml.ElementTree as ET
from pathlib import Path

import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
from docx import Document

from . import memory, render_pool, requisites
from .filehash import read_head
from .model import BBox, Line, Page, ParsedDoc, Word
from .normalize import nfc
from .ocr_ensemble import group_lines, run_ensemble

MIN_TEXT_CHARS = 20  # меньше — считаем, что текстового слоя нет
OCR_DPI = 300  # ТЗ 9.1: приёмка OCR на печатном тексте ≥ 300 dpi
LOW_QUALITY_CONF = (
    60.0  # средняя уверенность ансамбля ниже — LOW_QUALITY (OS-INSP-2.1.2)
)
REQ_DPI = 100  # растр для поиска печатей и подписей на страницах с текстовым слоем
_GRID = 100_000


# pdfium не потокобезопасен: одновременные вызовы из пула потоков FastAPI роняют процесс (SIGSEGV).
# Параллелизм в проде — процессами (uvicorn --workers N), внутри процесса — строго по одному.
PDFIUM_LOCK = threading.Lock()

# T-216: паспорта листов без текстового слоя, посчитанные разбором в режиме skip, — до записи в кэш (docstore). ParsedDoc
# не расширяется (PARSER_REV не трогаем): паспорт живёт отдельным файлом кэша passport-<sha>.
_PASSPORTS: dict[str, list[dict]] = {}
_PASSPORTS_LOCK = threading.Lock()


def take_passports(sha256: str) -> list[dict] | None:
    """Паспорта последнего разбора файла (и забыть их); None — разбор шёл без паспорта."""
    with _PASSPORTS_LOCK:
        return _PASSPORTS.pop(sha256, None)


memory.limit_arenas()  # до пулов страниц и потоков ONNX Runtime: каждая новая арена glibc держит свои освобождённые растры (T-230)


class UnsupportedFormat(ValueError):
    pass


class CorruptedFile(ValueError):
    pass


def detect_kind(path: Path) -> str:
    head = read_head(path, 8)  # сигнатура, а не весь файл (T-169: тома ИД до гигабайта)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        from .formats import is_xlsx  # GAP-INSP-04: XLSX сверх ТЗ (OS-INSP-1.2.9)

        return "xlsx" if is_xlsx(path) else "docx"
    from .formats import image_kind

    if image_kind(head):
        return "image"
    if head.lstrip().startswith(b"<") or head.startswith(b"\xef\xbb\xbf<"):
        return "xml"
    raise UnsupportedFormat(
        f"неподдерживаемый формат: {path.name}; поддерживаются PDF, DOCX, XML, XLSX, JPG, PNG, TIF"
    )


# ─────────────────────────────────────────────── PDF


def _to_norm(page: pdfium.PdfPage, x: float, y: float) -> tuple[float, float]:
    dx, dy = ctypes.c_int(), ctypes.c_int()
    pdfium_c.FPDF_PageToDevice(page.raw, 0, 0, _GRID, _GRID, 0, x, y, dx, dy)
    return min(max(dx.value / _GRID, 0.0), 1.0), min(max(dy.value / _GRID, 0.0), 1.0)


def norm_box(
    page: pdfium.PdfPage, left: float, bottom: float, right: float, top: float
) -> BBox:
    """Прямоугольник в координатах PDF → [x0, y0, x1, y1] в долях видимой страницы."""
    a = _to_norm(page, left, top)
    b = _to_norm(page, right, bottom)
    return (
        round(min(a[0], b[0]), 5),
        round(min(a[1], b[1]), 5),
        round(max(a[0], b[0]), 5),
        round(max(a[1], b[1]), 5),
    )


def union(boxes: list[BBox]) -> BBox | None:
    boxes = [b for b in boxes if b]
    if not boxes:
        return None
    return (
        min(b[0] for b in boxes),
        min(b[1] for b in boxes),
        max(b[2] for b in boxes),
        max(b[3] for b in boxes),
    )


# Переносы слов в текстовом слое: pdfium отдаёт маркер U+FFFE на месте дефиса переноса («конструктив\ufffeной»),
# встречаются мягкий дефис U+00AD и U+0002. Маркер выбрасывается — слово склеивается, якоря находят оборот
# (реальный лист ПБ, T-129: «класса конструктивной пожарной опасности здания С0» терялся целиком).
HYPHEN_MARKS = frozenset("\ufffe\u00ad\u0002")


def _raw_words(
    page: pdfium.PdfPage,
) -> list[tuple[str, tuple[float, float, float, float]]]:
    """Слова текстового слоя с рамкой в координатах PDF (до поворота страницы)."""
    tp = page.get_textpage()
    out: list[tuple[str, tuple[float, float, float, float]]] = []
    cur, boxes = "", []

    def flush() -> None:
        nonlocal cur, boxes
        if cur.strip() and boxes:
            out.append(
                (
                    cur,
                    (
                        min(b[0] for b in boxes),
                        min(b[1] for b in boxes),
                        max(b[2] for b in boxes),
                        max(b[3] for b in boxes),
                    ),
                )
            )
        cur, boxes = "", []

    for i in range(tp.count_chars()):
        ch = tp.get_text_range(i, 1)
        if ch in HYPHEN_MARKS:
            continue
        if ch.isspace():
            flush()
            continue
        l, b, r, t = tp.get_charbox(i)
        if r <= l or t <= b:
            continue
        # широкий разрыв (колонка таблицы) — новое слово; высота — по соседним символам:
        # у точки и запятой рамка крошечная, по ней порог ложно срабатывает («1.18» → «1 .18»)
        ref_h = max(t - b, boxes[-1][3] - boxes[-1][1]) if boxes else t - b
        if boxes and l - boxes[-1][2] > 1.5 * ref_h:
            flush()
        cur += ch
        boxes.append((l, b, r, t))
    flush()
    tp.close()
    return out


TALL = 2.5  # слово выше медианы страницы в 2,5 раза — вертикальный текст (поля штампа «Подп. и дата», «Инв. №»)


def group_rows(
    words: list[tuple[str, tuple[float, float, float, float]]],
) -> list[list[tuple[str, tuple[float, float, float, float]]]]:
    """Геометрические строки из слов (рамки в координатах PDF, ось y вверх).

    Слово входит в строку, если по вертикали перекрывается хотя бы на половину меньшей высоты со своим ближайшим по
    горизонтали соседом в этой строке: базовая линия непрерывна локально, а у соседних колонок она бывает сдвинута
    на треть строки (реальный лист АР, T-129). Раньше строка росла на высоту каждого добавленного слова: высокая
    рамка вертикального текста штампа «засасывала» соседние строки, и две строки абзаца перемешивались пословно
    («Степень Класс конструктивной огнестойкости пожарной…» — реальный лист ПОС). Высокие слова — отдельной строкой.
    """
    hs = sorted(b[3] - b[1] for _, b in words)
    h_med = hs[len(hs) // 2] if hs else 0.0
    words = sorted(words, key=lambda w: (-(w[1][1] + w[1][3]) / 2, w[1][0]))
    rows: list[list[tuple[str, tuple[float, float, float, float]]]] = []
    ext: list[
        list[float]
    ] = []  # [низ, верх, высокая?] — только предфильтр кандидатов, не критерий
    for w in words:
        (_, (l, b, r, t)) = w
        h = t - b
        if h_med and h > TALL * h_med:
            rows.append([w])
            ext.append([b, t, 1.0])
            continue
        best, best_ov = -1, 0.0
        for i, row in enumerate(rows):
            rb, rt, tall = ext[i]
            if tall or rt < b or rb > t:
                continue
            # ближайшее по горизонтали слово строки (зазор между рамками, 0 — перекрываются)
            nb = min(row, key=lambda z: max(z[1][0] - r, l - z[1][2], 0.0))[1]
            ov = min(t, nb[3]) - max(b, nb[1])
            if ov >= 0.5 * min(h, nb[3] - nb[1]) and ov > best_ov:
                best, best_ov = i, ov
        if best < 0:
            rows.append([w])
            ext.append([b, t, 0.0])
        else:
            rows[best].append(w)
            ext[best][0] = min(ext[best][0], b)
            ext[best][1] = max(ext[best][1], t)
    return rows


def _text_lines(page: pdfium.PdfPage) -> list[Line]:
    """Строки собираются геометрически (group_rows): pdfium отдаёт порядок чтения, но ячейки таблицы нередко
    рвёт на разные «строки»; для поиска «название … значение» важна именно геометрическая строка. Группировка
    идёт в координатах PDF до поворота — там текст горизонтален и на повёрнутых листах.
    """
    rows = group_rows(_raw_words(page))
    lines = []
    for row in rows:
        row.sort(key=lambda x: x[1][0])
        ws = [Word(text=t, bbox=norm_box(page, *box)) for t, box in row]
        lines.append(Line(text=" ".join(w.text for w in ws), words=ws))
    return lines


def ocr_workers(env: dict | None = None) -> int:
    """Страниц-сканов одного документа параллельно: INSPECTOR_OCR_WORKERS, по умолчанию 4, не меньше 1."""
    raw = (os.environ if env is None else env).get("INSPECTOR_OCR_WORKERS", "4")
    try:
        return max(1, int(raw))
    except ValueError:
        return 4


OCR_WORKERS = ocr_workers()


class PixelBudget:
    """Общий на процесс бюджет мегапикселей растров OCR в работе (T-129): растр листа А1 при 300 dpi — около 70 Мпикс
    и ~0,8 ГБ с копиями ансамбля; без бюджета 3 документа × 4 воркера убивали контейнер ML по памяти (OOM, стенд
    «Алтуфьево»). Страница занимает бюджет по своему размеру; страница больше бюджета идёт одна, а не ждёт вечно."""

    def __init__(self, mpx: float):
        self.limit = mpx
        self.used = 0.0
        self.cv = threading.Condition()

    def acquire(self, mpx: float) -> float:
        need = min(mpx, self.limit)
        with self.cv:
            self.cv.wait_for(lambda: self.used + need <= self.limit or self.used == 0)
            self.used += need
        return need

    def release(self, need: float) -> None:
        with self.cv:
            self.used -= need
            self.cv.notify_all()


def _budget_mpx(env: dict | None = None) -> float:
    raw = (os.environ if env is None else env).get("INSPECTOR_OCR_BUDGET_MPX", "150")
    try:
        return max(1.0, float(raw))
    except ValueError:
        return 150.0


OCR_BUDGET = PixelBudget(_budget_mpx())


def page_mpx(width_pt: float, height_pt: float, dpi: int = OCR_DPI) -> float:
    """Мегапиксели растра страницы при рендере в dpi (размеры PDF — в пунктах, 1/72 дюйма)."""
    return (width_pt / 72 * dpi) * (height_pt / 72 * dpi) / 1e6


def _env_num(name: str, default: float, lo: float) -> float:
    try:
        return max(lo, float(os.environ.get(name, default)))
    except ValueError:
        return default


# SEC-01 (OWASP LLM10/API4, аудит T-129): размер растра задаёт недоверенный PDF (MediaBox). Потолок на страницу: лист
# А0 при 300 dpi — ~140 Мпикс и проходит как есть; лист больше рендерится с пониженным dpi и помечается LOW_QUALITY.
# Потолок числа страниц: 50 МБ пустых «сканов» иначе — часы OCR.
MAX_PAGE_MPX = _env_num("INSPECTOR_MAX_PAGE_MPX", 150.0, 1.0)
MAX_PAGES = int(_env_num("INSPECTOR_MAX_PAGES", 2000, 1))


class TooLarge(CorruptedFile):
    """Документ вне пределов разбора (число страниц): отказ 422 с причиной, а не часы работы."""


def render_scale(
    width_pt: float, height_pt: float, dpi: int = OCR_DPI, max_mpx: float | None = None
) -> float:
    """Масштаб рендера pdfium: dpi/72, но не больше потолка мегапикселей на страницу."""
    cap = MAX_PAGE_MPX if max_mpx is None else max_mpx
    want = dpi / 72
    from .resource_scope import require_render_pixels
    # Admission never lowers requested DPI to make a reservation fit. The
    # existing page safety ceiling remains a separate parser policy.
    require_render_pixels(width_pt, height_pt, want)
    if width_pt <= 0 or height_pt <= 0:
        return want
    return min(want, (cap * 1e6 / (width_pt * height_pt)) ** 0.5)


def _downscaled(width_pt: float, height_pt: float) -> bool:
    return render_scale(width_pt, height_pt) < OCR_DPI / 72 - 1e-9


def _ocr_image(
    img, number: int, w: float, h: float, rot: int, downscaled: bool = False
) -> Page:
    """Страница без текстового слоя по готовому растру: ансамбль OCR с голосованием (OS-INSP-2.1.5–2.1.7).
    pdfium не трогает — выполняется вне PDFIUM_LOCK."""
    res = run_ensemble(img)
    if (
        not res.engines
    ):  # ни одного доступного движка — воздерживаемся, а не отдаём пустой «OK»
        return Page(
            page=number,
            width=w,
            height=h,
            rotation=rot,
            source="ocr",
            quality="ABSTAIN",
            lines=[],
            requisites=requisites.detect(img, []),
        )
    lines = (
        res.lines if res.lines is not None else group_lines(res.words)
    )  # повёрнутый скан — строки в кадре прямого текста
    mean = res.mean_conf
    # лист больше потолка разобран с пониженным dpi — качество не гарантируем (SEC-01)
    quality = (
        "OK" if lines and mean >= LOW_QUALITY_CONF and not downscaled else "LOW_QUALITY"
    )
    return Page(
        page=number,
        width=w,
        height=h,
        rotation=rot,
        source="ocr",
        quality=quality,
        ocr_confidence=round(mean, 1),
        lines=lines,
        engines=res.engines,
        execution_failures=res.execution_failures,
        anchor_words=res.anchor_words,
        disputed_words=res.disputed_words,
        agreement=res.agreement,
        requisites=requisites.detect(img, lines),
    )


def _ocr_page(page: pdfium.PdfPage, number: int) -> Page:
    """Совместимость: страница без текстового слоя целиком под замком (растр + OCR). Конвейер идёт через _parse_pdf."""
    w, h = page.get_size()
    from .resource_scope import render_guarded
    scale = render_scale(w, h)
    return _ocr_image(
        render_guarded(lambda: page.render(scale=scale).to_pil(), w, h, scale),
        number,
        w,
        h,
        page.get_rotation(),
        downscaled=_downscaled(w, h),
    )


def parse_pdf(path: Path, sha256: str) -> ParsedDoc:
    """Замок pdfium берётся только на обращения к pdfium (текст, рендер), а не на весь документ: OCR страниц-сканов
    идёт вне замка и параллельно (OCR_WORKERS). Раньше замок держался и на OCR — документ со 110 сканами
    останавливал разбор всех остальных (стенд T-129, «Алтуфьевское 79Б»)."""
    return _parse_pdf(path, sha256)


def pdf_pages(path: Path) -> int:
    """Число страниц PDF — только оглавление документа, без текста и растров."""
    with PDFIUM_LOCK:
        try:
            doc = pdfium.PdfDocument(str(path))
        except pdfium.PdfiumError as e:
            raise CorruptedFile(f"повреждённый PDF: {path.name}") from e
        try:
            return len(doc)
        finally:
            doc.close()


def parse_pdf_part(path: Path, sha256: str, first: int, last: int) -> ParsedDoc:
    """T-233: часть тома — страницы [first, last) с номерами как в целом документе. Части одного тома разбираются
    разными процессами ML параллельно и собираются docstore.assemble_parts; предел MAX_PAGES — на часть, не на том."""
    return _parse_pdf(path, sha256, first, last)


def _parse_pdf(
    path: Path, sha256: str, first: int = 0, last: int | None = None
) -> ParsedDoc:
    from .resource_scope import render_guarded, require_single_ocr_worker
    require_single_ocr_worker(OCR_WORKERS, render_pool.procs())
    with PDFIUM_LOCK:
        try:
            doc = pdfium.PdfDocument(str(path))
        except pdfium.PdfiumError as e:
            raise CorruptedFile(f"повреждённый PDF: {path.name}") from e
        n = len(doc)
        stop = n if last is None else min(last, n)
        if first < 0 or first >= stop:
            doc.close()
            raise ValueError(
                f"{path.name}: диапазон страниц [{first}, {last}) вне документа из {n} стр."
            )
        if stop - first > MAX_PAGES:
            doc.close()
            raise TooLarge(
                f"{path.name}: {n} страниц — больше предела разбора {MAX_PAGES} (INSPECTOR_MAX_PAGES)"
            )
    from .page_passport import SKIP_CLASSES, mode as passport_mode, passport

    skip_mode = passport_mode() == "skip"
    try:
        pages: dict[int, Page] = {}
        scans: list[int] = []
        passports: list[dict] = []
        for i in range(first, stop):
            with PDFIUM_LOCK:
                page = doc[i]
                lines = _text_lines(page)
                if sum(len(ln.text) for ln in lines) >= MIN_TEXT_CHARS:
                    w, h = page.get_size()
                    rot = page.get_rotation()
                    # печати и подписи — растровые даже на листе с текстовым слоем
                    scale = render_scale(w, h, REQ_DPI)
                    req_img = render_guarded(lambda: page.render(scale=scale).to_pil(), w, h, scale)
                else:
                    req_img = None
                    # T-216: паспорт только для листов без текстового слоя — кандидатов в OCR; пустой лист и лист
                    # подписи не рисуются и не идут в OCR, но остаются в документе (номера страниц не сдвигаются)
                    pp = passport(page) if skip_mode else None
                    if pp is not None:
                        pp["page"] = i + 1
                        passports.append(pp)
                    if pp is not None and pp["page_class"] in SKIP_CLASSES:
                        w, h = page.get_size()
                        pages[i] = Page(
                            page=i + 1,
                            width=w,
                            height=h,
                            rotation=page.get_rotation(),
                            source="skipped",
                            quality="OK",
                            lines=[],
                        )
                    else:
                        scans.append(i)
                page.close()
            if req_img is not None:
                pages[i] = Page(
                    page=i + 1,
                    width=w,
                    height=h,
                    rotation=rot,
                    source="text",
                    lines=lines,
                    requisites=requisites.detect(req_img, lines),
                )

        def ocr(i: int) -> Page:
            with PDFIUM_LOCK:
                page = doc[i]
                w, h, rot = *page.get_size(), page.get_rotation()
                page.close()
            scale = render_scale(w, h)
            held = OCR_BUDGET.acquire(
                page_mpx(w, h, scale * 72)
            )  # память: не больше бюджета пикселей на процесс
            try:
                # T-233: страница рисуется в пуле процессов — без замка процесса, параллельно; пул выключен или лист
                # крупный — под замком, как раньше
                img = render_pool.render(
                    str(path), i, scale, page_mpx(w, h, scale * 72)
                )
                if img is None:
                    with PDFIUM_LOCK:  # растр готовится под замком и сразу; копятся не растры, а готовые страницы
                        page = doc[i]
                        img = render_guarded(lambda: page.render(scale=scale).to_pil(), w, h, scale)
                        page.close()
                return _ocr_image(img, i + 1, w, h, rot, downscaled=_downscaled(w, h))
            finally:
                OCR_BUDGET.release(held)

        if scans:
            with ThreadPoolExecutor(max_workers=min(OCR_WORKERS, len(scans))) as pool:
                from .execution_scope import scoped_submit
                futures = [scoped_submit(pool, ocr, i) for i in scans]
                for i, future in zip(scans, futures):
                    pages[i] = future.result()
    finally:
        with PDFIUM_LOCK:
            doc.close()
        memory.release()  # растры документа — обратно системе, а не в арены потоков (T-230)
    ordered = [pages[i] for i in range(first, stop)]
    engines = {"pdfium"} | {e for p in ordered for e in p.engines}
    if skip_mode:
        with _PASSPORTS_LOCK:
            _PASSPORTS[sha256] = passports
    return ParsedDoc(
        sha256=sha256, kind="pdf", pages=ordered, engine="+".join(sorted(engines))
    )


# ─────────────────────────────────────────────── DOCX, XML


def _plain_lines(texts: list[str]) -> list[Line]:
    out = []
    for t in texts:
        t = nfc(t)
        if t:
            out.append(Line(text=t, words=[Word(text=w) for w in t.split(" ")]))
    return out


def parse_docx(path: Path, sha256: str) -> ParsedDoc:
    try:
        d = Document(str(path))
    except Exception as e:  # python-docx бросает разнородные исключения на битом архиве
        raise CorruptedFile(f"повреждённый DOCX: {path.name}") from e
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            texts.append("  ".join(c.text for c in row.cells))
    return ParsedDoc(
        sha256=sha256,
        kind="docx",
        engine="python-docx",
        pages=[
            Page(
                page=1,
                width=0,
                height=0,
                source="structured",
                lines=(lines := _plain_lines(texts)),
                requisites=requisites.detect(None, lines),
            )
        ],
    )


def parse_xml(path: Path, sha256: str) -> ParsedDoc:
    try:
        root = ET.parse(str(path)).getroot()
    except (
        ET.ParseError,
        ValueError,
    ) as e:  # ValueError — DefusedXmlException (DTD, сущности)
        raise CorruptedFile(f"повреждённый XML: {path.name}") from e
    texts = []
    for el in root.iter():
        label = el.attrib.get("наименование") or el.attrib.get("name")
        text = (el.text or "").strip()
        if label and text:
            texts.append(f"{label} {el.attrib.get('ед', '')} {text}")
        elif text:
            texts.append(text)
    return ParsedDoc(
        sha256=sha256,
        kind="xml",
        engine="etree",
        pages=[
            Page(
                page=1,
                width=0,
                height=0,
                source="structured",
                lines=(lines := _plain_lines(texts)),
                requisites=requisites.detect(None, lines),
            )
        ],
    )


def page_raster(path: Path, doc: ParsedDoc, number: int):
    """Растр OCR-страницы тем же способом, что при разборе (OCR_DPI): рамки слов — доли именно этого растра.
    Нужен для точечного перечитывания значения (reread, OS-INSP-2.2.12). Не OCR-страница или нет файла — None."""
    page = next((p for p in doc.pages if p.page == number), None)
    if page is None or page.source != "ocr" or not path.is_file():
        return None
    if doc.kind == "pdf":
        with PDFIUM_LOCK:
            pdf = pdfium.PdfDocument(str(path))
            try:
                if number > len(pdf):
                    return None
                pg = pdf[number - 1]
                w, h = pg.get_size()
                from .resource_scope import render_guarded
                scale = render_scale(w, h)
                return render_guarded(lambda: pg.render(scale=scale).to_pil(), w, h, scale)
            finally:
                pdf.close()
    if doc.kind == "image":
        from .formats import image_page

        return image_page(path, number, OCR_DPI)
    return None


def parse_file(path: Path, sha256: str) -> ParsedDoc:
    kind = detect_kind(path)
    from .formats import parse_image, parse_xlsx

    return {
        "pdf": parse_pdf,
        "docx": parse_docx,
        "xml": parse_xml,
        "xlsx": parse_xlsx,
        "image": parse_image,
    }[kind](path, sha256)
