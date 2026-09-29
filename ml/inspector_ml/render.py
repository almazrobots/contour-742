"""Рендер протокола в PDF и DOCX (OS-INSP-5.1). На вход — JSON протокола из API.

Весь текст из данных идёт через safe_markup: в PDF — pdf_text (разметка не исполняется, OS-INSP-5.1.20),
в DOCX — docx_text (без недопустимых в XML символов, OS-INSP-5.1.21)."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from .paths import repo_root
from .safe_markup import docx_text, pdf_text

FONT = repo_root() / "assets/fonts/NotoSans.ttf"
if "Noto" not in pdfmetrics.getRegisteredFontNames():
    pdfmetrics.registerFont(TTFont("Noto", str(FONT)))

STATUS_RU = {
    "CANDIDATE": "Кандидат",
    "CONFIRMED_VIOLATION": "Подтверждённое нарушение",
    "NEGATIVE_VERIFIED": "Расхождение не подтверждено",
    "MISSING_EVIDENCE": "Нет доказательства",
    "NOT_APPLICABLE": "Неприменимо",
    "NOT_COMPARABLE": "Несопоставимо",
    "CLARIFICATION_REQUIRED": "Требует уточнения",
}


def _sources(card: dict) -> str:
    return "; ".join(
        f"{s['stage']} {s['document_code']} ред.{s['revision']} стр.{s['page']}"
        + (
            " [" + ", ".join(f"{x:.3f}" for x in s["bbox"]) + "]"
            if s.get("bbox")
            else ""
        )
        for s in card.get("sources", [])
    )


def _rows(p: dict) -> list[tuple[str, list[list[str]]]]:
    s = p["sections"]
    return [
        (
            "1. Комплектность и сопоставимость",
            [["Код", "Параметр", "Статус", "Основание"]]
            + [
                [
                    r["param_code"],
                    r["parameter_name"],
                    STATUS_RU.get(r["status"], r["status"]),
                    r.get("reason") or "",
                ]
                for r in s["completeness"]
            ],
        ),
        (
            "2. Предварительные кандидаты",
            [["Finding", "Параметр", "Ожидается", "Факт", "Δ", "Источники"]]
            + [
                [
                    c["finding_id"],
                    c["parameter_name"],
                    c.get("expected") or "",
                    c.get("actual") or "",
                    c.get("delta") or "",
                    _sources(c),
                ]
                for c in s["candidates"]
            ],
        ),
        (
            "3. Подтверждённые инспектором нарушения",
            [["Finding", "Параметр", "Ожидается", "Факт", "Решение", "Источники"]]
            + [
                [
                    c["finding_id"],
                    c["parameter_name"],
                    c.get("expected") or "",
                    c.get("actual") or "",
                    ((c.get("decision") or {}).get("user_name") or "")
                    + " "
                    + ((c.get("decision") or {}).get("comment") or ""),
                    _sources(c),
                ]
                for c in s["confirmed_violations"]
            ],
        ),
        (
            "4. Проверенные отрицательные результаты",
            [["Код", "Параметр", "Ожидается", "Факт", "Кем"]]
            + [
                [
                    r["param_code"],
                    r["parameter_name"],
                    r.get("expected") or "",
                    r.get("actual") or "",
                    "инспектор" if r.get("by") == "inspector" else "система",
                ]
                for r in s["negative_verified"]
            ],
        ),
        (
            "5. Гипотезы свободного поиска (не нарушения)",
            [["Метод", "Описание", "Норматив", "Приоритет"]]
            + [
                [
                    h["discovery_method"],
                    h["description"],
                    h.get("normative_base") or "",
                    h.get("review_priority") or "",
                ]
                for h in s["suspicions"]
            ],
        ),
    ]


def _header(p: dict) -> list[tuple[str, str]]:
    v = p["versions"]
    return [
        ("Объект", ", ".join(x for x in (p["object"]["name"], p["object"].get("address")) if x)),  # без адреса — без висячей запятой
        ("Разрешение", p["object"].get("permit_number") or "—"),
        ("Проверка (process_id)", p["process_id"]),
        ("Версия протокола", str(p["protocol_version"])),
        ("Статус", p["status"]),
        (
            "Тип проверки",
            f"{p['check_type'].get('scenario')} — {p['check_type'].get('title')}",
        ),
        ("Статус загрузки", ", ".join(p["upload_status"])),
        (
            "Версии",
            f"Матрица {v['matrix_version']} · модель {v['model_version']} · набор {v['dataset_version']}",
        ),
        ("Хеш реестра входных файлов", v["input_manifest_hash"]),
    ]


# узкие колонки таблиц разделов; остальные делят ширину поровну. Явные ширины: при автоширине длинное
# неразрывное значение из документа (путь, хеш, URL) отбирало у соседей всю ширину и роняло выгрузку (OS-INSP-5.1.21)
NARROW = {"Код", "Finding", "Статус", "Δ", "Кем", "Метод", "Приоритет"}


def _col_widths(head: list[str], total: float) -> list[float]:
    weights = [1 if h in NARROW else 2 for h in head]
    return [total * w / sum(weights) for w in weights]


def render_pdf(p: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"Протокол {p['process_id']} v{p['protocol_version']}",
    )
    h1 = ParagraphStyle("h1", fontName="Noto", fontSize=15, leading=19, spaceAfter=6)
    h2 = ParagraphStyle(
        "h2",
        fontName="Noto",
        fontSize=11.5,
        leading=15,
        spaceBefore=10,
        spaceAfter=4,
        textColor=colors.HexColor("#4b2fd4"),
    )
    cell = ParagraphStyle("c", fontName="Noto", fontSize=7.5, leading=9.5)
    story = [Paragraph("Протокол камеральной проверки ПД / РД / ИД", h1)]
    story.append(
        Table(
            [[Paragraph(pdf_text(k), cell), Paragraph(pdf_text(v), cell)] for k, v in _header(p)],
            colWidths=[55 * mm, 210 * mm],
            style=[
                ("FONT", (0, 0), (-1, -1), "Noto", 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d9dbe6")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f4f3ff")),
            ],
        )
    )
    for title, rows in _rows(p):
        story.append(Paragraph(pdf_text(f"{title} — {len(rows) - 1}"), h2))
        if len(rows) == 1:
            story.append(Paragraph("Нет записей", cell))
            continue
        t = Table(
            [[Paragraph(pdf_text(c), cell) for c in r] for r in rows],
            colWidths=_col_widths(rows[0], doc.width),
            repeatRows=1,
        )
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d9dbe6")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efeefe")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(t)
    story.append(Spacer(1, 8 * mm))
    story.append(
        Paragraph(
            "Уровень риска определяет только очередность экспертной проверки и не является основанием для предписания. Нарушением признаётся только запись, подтверждённая инспектором.",
            cell,
        )
    )
    doc.build(story)
    return buf.getvalue()


def render_docx(p: dict) -> bytes:
    d = Document()
    d.styles["Normal"].font.size = Pt(9)
    d.add_heading("Протокол камеральной проверки ПД / РД / ИД", level=1)
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in _header(p):
        r = t.add_row().cells
        r[0].text, r[1].text = docx_text(k), docx_text(v)
    for title, rows in _rows(p):
        d.add_heading(f"{title} — {len(rows) - 1}", level=2)
        if len(rows) == 1:
            d.add_paragraph("Нет записей")
            continue
        tb = d.add_table(rows=0, cols=len(rows[0]))
        tb.style = "Table Grid"
        for row in rows:
            cells = tb.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = docx_text(val)
    d.add_paragraph(
        "Уровень риска определяет только очередность экспертной проверки. Нарушением признаётся только запись, подтверждённая инспектором."
    )
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
