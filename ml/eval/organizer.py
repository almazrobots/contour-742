"""Адаптер разметки организаторов → эталонные evidence_group стенда.

Источник — лист «ПРИМЕРЫ РАЗМЕТКИ» в ТЗ/Матрица_параметров_редакция1.1.xlsx (вне git). Колонка
«Источники / страницы / bbox [0;1]» имеет вид

    PD:ALT79B-000015:стр.19:bbox [0.78,0.10,0.91,0.35];[0.36,0.10,0.64,0.25] | RD:ALT79B-000077:стр.4:bbox […]

— фрагменты через «|», в фрагменте стадия, file_id, страница и ноль или больше рамок через «;»
(«bbox —» — рамки нет). Сами документы у нас отсутствуют: адаптер только переводит формат, чтобы
скрытый тест в той же разметке сразу лёг в стенд (метка NEGATIVE_VERIFIED → FPR, CONFIRMED_VIOLATION
→ положительная группа).
"""

from __future__ import annotations

import re
from pathlib import Path

SHEET = "ПРИМЕРЫ РАЗМЕТКИ"
_FRAG = re.compile(
    r"^\s*(PD|RD|ID)\s*:\s*([^:|]+?)\s*:\s*стр\.?\s*(\d+)\s*:\s*bbox\s*(.*?)\s*$", re.IGNORECASE
)
_BOX = re.compile(r"\[\s*([-\d.,\s]+?)\s*\]")


class SourceFormatError(ValueError):
    pass


def parse_bboxes(s: str) -> list[list[float]]:
    s = s.strip()
    if s in ("", "—", "-", "–"):
        return []
    boxes = []
    for m in _BOX.finditer(s):
        nums = [float(x) for x in re.split(r"\s*,\s*", m.group(1).strip()) if x]
        if (
            len(nums) != 4
            or not all(0.0 <= v <= 1.0 for v in nums)
            or nums[0] > nums[2]
            or nums[1] > nums[3]
        ):
            raise SourceFormatError(
                f"рамка не [x0,y0,x1,y1] в долях [0;1]: {m.group(0)}"
            )
        boxes.append(nums)
    if not boxes:
        raise SourceFormatError(f"не разобрана рамка: {s!r}")
    return boxes


def parse_sources(s: str) -> list[dict]:
    """Строка источников → [{stage, file_id, page, bboxes}] в порядке записи."""
    out = []
    for part in (s or "").split("|"):
        if not part.strip():
            continue
        m = _FRAG.match(part)
        if not m:
            raise SourceFormatError(
                f"фрагмент не по формату «СТАДИЯ:file_id:стр.N:bbox …»: {part.strip()!r}"
            )
        out.append(
            {
                "stage": m.group(1).upper(),
                "file_id": m.group(2),
                "page": int(m.group(3)),
                "bboxes": parse_bboxes(m.group(4)),
            }
        )
    return out


def evidence(sources: list[dict]) -> list[dict]:
    """Фрагменты → элементы evidence стенда: по одному на рамку; без рамки — bbox None."""
    out = []
    for s in sources:
        for b in s["bboxes"] or [None]:
            out.append(
                {
                    "file_id": s["file_id"],
                    "page": s["page"],
                    "stage": s["stage"],
                    "bbox": b,
                }
            )
    return out


def to_group(row: dict) -> dict:
    src = parse_sources(row["sources"])
    fid = row["finding_id"]
    return {
        "evidence_group_id": fid,
        # object_id — префикс file_id (ALT79B-000015 → ALT79B): так объект изолируется в бутстрапе
        "object_id": src[0]["file_id"].split("-")[0] if src else fid.split("-")[0],
        "object_name": row.get("object"),
        "param": None,  # в примерах параметр Матрицы не указан — сопоставление по finding_id
        "label": row["status"],
        "expected_value": row.get("expected"),
        "actual_value": row.get("actual"),
        "expert_action": row.get("action"),
        "label_rule": row.get("rule"),
        "evidence": evidence(src),
    }


def load_examples(xlsx: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(xlsx), read_only=True, data_only=True)
    ws = wb[SHEET]
    groups, header = [], None
    for row in ws.iter_rows(values_only=True):
        cells = [("" if c is None else str(c)).strip() for c in row]
        if header is None:
            if cells and cells[0] == "finding_id":
                header = cells
            continue
        if not cells or not cells[0]:
            continue
        rec = dict(
            zip(
                [
                    "finding_id",
                    "status",
                    "object",
                    "sources",
                    "expected",
                    "actual",
                    "action",
                    "rule",
                ],
                cells,
            )
        )
        groups.append(to_group(rec))
    wb.close()
    if header is None:
        raise SourceFormatError(f"на листе «{SHEET}» нет строки заголовка finding_id")
    return groups


def find_xlsx(start: Path | None = None) -> Path | None:
    """ТЗ/ лежит в корне основного checkout (вне git) — ищем вверх от ml/, worktree тоже найдёт."""
    here = (start or Path(__file__)).resolve()
    for p in here.parents:
        f = p / "ТЗ" / "Матрица_параметров_редакция1.1.xlsx"
        if f.exists():
            return f
    return None
