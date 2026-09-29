"""Faithful, immutable structured previews from verified original bytes, never OCR.
Private JSON is rendered as text/table by the authenticated annotation viewer.
"""
import datetime as dt
import hashlib
import io
import json
import os
import re
import unicodedata
from functools import lru_cache
from pathlib import Path

REVISION = 'structured-original.v1'

def norm(text):
    return ' '.join(unicodedata.normalize('NFC', str(text)).split())

def cell(value):
    if value is None:return ''
    if isinstance(value,bool):return 'да' if value else 'нет'
    if isinstance(value,float):return str(int(value)) if value.is_integer() else repr(value).replace('.',',')
    if isinstance(value,(dt.date,dt.datetime)):return value.strftime('%d.%m.%Y')
    return unicodedata.normalize('NFC',str(value))

@lru_cache(maxsize=4)
def original(path,sha,kind):
    data=Path(path).read_bytes()
    if hashlib.sha256(data).hexdigest()!=sha:raise ValueError('original_sha_mismatch')
    if kind=='xlsx':
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        try:
            return [{'title':sheet.title,'rows':[[cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]} for sheet in book.worksheets]
        finally:book.close()
    if kind=='docx':
        from docx import Document
        doc=Document(io.BytesIO(data))
        lines=[p.text for p in doc.paragraphs]
        lines.extend('  '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
    elif kind=='xml':
        from defusedxml.ElementTree import fromstring
        lines=[]
        for node in fromstring(data).iter():
            label=node.attrib.get('наименование') or node.attrib.get('name')
            text=(node.text or '').strip()
            if text:lines.append(f"{label} {node.attrib.get('ед','')} {text}" if label else text)
    else:raise ValueError('unsupported_structured_kind')
    return [{'title':kind.upper(),'rows':[[unicodedata.normalize('NFC',line)] for line in lines if line]}]

def preview(candidate,kind,staging,corpus='/opt/corpus/blobs'):
    sha=candidate['source_sha256']
    if not re.fullmatch('[a-f0-9]{64}',sha):raise ValueError('invalid_sha')
    page=int(candidate['page']);sheets=original(str(Path(corpus)/sha),sha,kind)
    if page<1 or page>len(sheets):raise ValueError('missing_structured_page')
    sheet=sheets[page-1];rows=sheet['rows'];needle=norm(candidate['extraction'].get('line_text',''))
    if not needle:raise ValueError('missing_source_line')
    # Some extractors quote a sentence or a window spanning adjacent original rows.
    # Locate the exact quote in the normalized original stream, retaining row offsets.
    text='';offsets=[]
    for i,row in enumerate(rows):
        line=norm(' '.join(v for v in row if v))
        if not line:continue
        if text:text+=' '
        offsets.append((i,len(text),len(text)+len(line)));text+=line
    positions=[match.start() for match in re.finditer('(?='+re.escape(needle)+')',text)]
    if len(positions)!=1:raise ValueError('ambiguous_source_line' if positions else 'source_line_not_found')
    start=positions[0];end=start+len(needle)
    targets=[i for i,lo,hi in offsets if lo<end and hi>start]
    target=targets[0];first=max(0,target-3);last=min(len(rows),targets[-1]+4)
    if last-first>40:raise ValueError('preview_too_many_rows')
    context=rows[first:last]
    if any(len(row)>256 for row in context):raise ValueError('preview_too_wide')
    label=f"XLSX · лист «{sheet['title']}» · строки {first+1}–{last}" if kind=='xlsx' else f"{kind.upper()} · текстовое представление · строки {first+1}–{last}"
    document={'schema':REVISION,'kind':kind,'source_sha256':sha,'sheet':page,'label':label,
              'note':'Значения ячеек оригинала; формулы показаны по сохранённым результатам, без пересчёта. Оформление книги не воспроизводится.' if kind=='xlsx' else 'Текст из оригинального файла. Печатная вёрстка и номера страниц не воспроизводятся.',
              'first_row':first+1,'target_row':target+1,'target_rows':[i+1 for i in targets],'rows':context}
    raw=json.dumps(document,ensure_ascii=False,separators=(',',':')).encode()
    if len(raw)>256000:raise ValueError('preview_too_large')
    digest=hashlib.sha256(raw).hexdigest();root=Path(staging)/'structured';root.mkdir(mode=0o700,exist_ok=True)
    if os.geteuid()==0:os.chown(root,-1,65532);root.chmod(0o710)
    path=root/(digest+'.json')
    try:
        with path.open('xb') as out:out.write(raw)
    except FileExistsError:
        if path.read_bytes()!=raw:raise ValueError('preview_changed')
    path.chmod(0o640)
    if os.geteuid()==0:os.chown(path,-1,65532)
    return {'key':digest,'sha256':digest,'label':label,'revision':REVISION}
