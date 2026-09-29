import hashlib,importlib.util,json,tempfile,unittest
from pathlib import Path
spec=importlib.util.spec_from_file_location('structured',Path(__file__).parents[1]/'runner'/'verification-structured.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
class StructuredOriginal(unittest.TestCase):
    def make(self,root,path,kind,line,page=1):
        sha=hashlib.sha256(path.read_bytes()).hexdigest();(root/sha).write_bytes(path.read_bytes())
        c={'source_sha256':sha,'page':page,'extraction':{'line_text':line}}
        descriptor=m.preview(c,kind,root,root)
        data=(root/'structured'/(descriptor['key']+'.json')).read_bytes()
        self.assertEqual(hashlib.sha256(data).hexdigest(),descriptor['sha256'])
        self.assertEqual(m.preview(c,kind,root,root),descriptor)
        return json.loads(data),c
    def test_excel_exact_cells_sheet_row_and_empty_cells(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);path=root/'test.xlsx';book=Workbook();book.active.title='Лист 1';book.active.append(['Контекст']);book.active.append(['Площадь',None,123.5,'м²']);book.save(path)
            view,c=self.make(root,path,'xlsx','Площадь 123,5 м²')
            self.assertEqual(view['target_row'],2);self.assertEqual(view['rows'][1],['Площадь','', '123,5','м²']);self.assertIn('сохранённым',view['note'])
            c['extraction']['line_text']='нет такой строки'
            with self.assertRaisesRegex(ValueError,'source_line_not_found'):m.preview(c,'xlsx',root,root)
    def test_duplicate_rows_are_not_a_guessed_location(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);book=Workbook();book.active.append(['Этажи',11]);book.active.append(['Этажи',11]);path=root/'test.xlsx';book.save(path)
            sha=hashlib.sha256(path.read_bytes()).hexdigest();(root/sha).write_bytes(path.read_bytes())
            with self.assertRaisesRegex(ValueError,'ambiguous_source_line'):m.preview({'source_sha256':sha,'page':1,'extraction':{'line_text':'Этажи 11'}},'xlsx',root,root)
    def test_docx_paragraph_and_table_original_text(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);doc=Document();doc.add_paragraph('Проект');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Этажей';table.cell(0,1).text='11';path=root/'test.docx';doc.save(path)
            view,c=self.make(root,path,'docx','Этажей  11');self.assertEqual(view['rows'][1],['Этажей  11']);self.assertIn('вёрстка',view['note'])
    def test_exact_sentence_can_span_rows_without_inventing_a_single_row(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);doc=Document();doc.add_paragraph('Проектируемое здание');doc.add_paragraph('имеет 11 этажей.');path=root/'test.docx';doc.save(path)
            view,c=self.make(root,path,'docx','здание имеет 11 этажей.')
            self.assertEqual(view['target_rows'],[1,2]);self.assertEqual(view['rows'],[['Проектируемое здание'],['имеет 11 этажей.']])
    def test_xml_is_text_not_html_and_external_entities_refused(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);path=root/'test.xml';path.write_text('<root><value name="Площадь" ед="м2">123</value><x>&lt;script&gt;</x></root>')
            view,c=self.make(root,path,'xml','Площадь м2 123');self.assertEqual(view['rows'][1],['<script>'])
            evil=b'<!DOCTYPE x [<!ENTITY a SYSTEM "file:///etc/passwd">]><x>&a;</x>';sha=hashlib.sha256(evil).hexdigest();(root/sha).write_bytes(evil)
            with self.assertRaises(Exception):m.preview({**c,'source_sha256':sha},'xml',root,root)
    def test_sha_and_path_guard(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);(root/('a'*64)).write_text('wrong bytes')
            with self.assertRaisesRegex(ValueError,'original_sha_mismatch'):m.original(str(root/('a'*64)),'a'*64,'xml')
            with self.assertRaisesRegex(ValueError,'invalid_sha'):m.preview({'source_sha256':'../../etc/passwd'},'xml',root,root)
