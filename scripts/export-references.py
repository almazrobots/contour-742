#!/usr/bin/env python3
"""Выгрузка справочников Инспектора ИИ в Excel для сверки владельцем (T-233).

Источник — сидовые файлы data/seed/*.json и паспорта параметров data/seed/passports{,/draft,-pending}/M-*.json.
Результат — docs/справочники/справочники.xlsx: лист «Оглавление» (что откуда) и по листу на каждую таблицу справочника.
На каждом листе справа пустая колонка «Проверка владельца» — для отметки «верно / неверно» и комментария.

Запуск (на раннере, не на маке):
    scripts/remote-run.sh --light --get docs/справочники "cd ml && uv run python ../scripts/export-references.py"
Справочники не правятся этим скриптом: он только читает и раскладывает по таблицам.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SEED = ROOT / "data/seed"
OUT = ROOT / "docs/справочники"

# файл → (что это, как называем листы). Тексты ТЗ (tz-*.json) и схемы — не справочники, в выгрузку не входят.
FILES = [
    (
        "matrix.json",
        "Матрица параметров — 132 параметра проверки: раздел, источники в ПД/РД/ИД, логика срабатывания",
    ),
    (
        "norms.json",
        "Нормативная база для поиска норм к гипотезам (документы, названные в ТЗ и Матрице)",
    ),
    (
        "legal-acts.json",
        "Перечень нормативных правовых актов из ТЗ §2 «Нормативная база»",
    ),
    (
        "analogs.json",
        "Канон материалов и марок и таблица аналогов для сравнения замен (CMP-05, CMP-21)",
    ),
    (
        "scales.json",
        "Шкалы классов (огнестойкость, пожарная опасность и др.): значения от худшего к лучшему",
    ),
    ("catalog-ops.json", "Каталог операций сравнения TO-BE (слои и операции)"),
    ("param-deps.json", "Зависимости параметров: причина → следствие (каскады)"),
    ("w1-wave.json", "Реестр волны W1: параметр → операторы сравнения"),
    (
        "w2-wave.json",
        "Реестр волны W2: параметр → операторы → путь (текст/сущность/геометрия)",
    ),
    (
        "w2-text-kinds.json",
        "Текстовый путь W2: что проверяется на существующих видах паспорта",
    ),
    ("acceptance-metrics.json", "Приёмочные метрики ТЗ §14.3 и их пороги"),
    (
        "l8-gates.json",
        "Пороги ворот и проверок уровня L8 (уверенность извлечения и др.)",
    ),
    ("geom-gates.json", "Допуски проверок геометрии чертежей"),
    (
        "class-fold-vectors.json",
        "Контрольные пары нормализации обозначения класса (что во что превращается)",
    ),
]
PASSPORT_DIRS = [
    ("passports", "действует"),
    ("passports/draft", "черновик (draft)"),
    ("passports-pending", "ожидает (pending)"),
]
CHECK_COL = "Проверка владельца (верно / неверно / комментарий)"
HEAD_FILL = PatternFill("solid", fgColor="DDE6F3")
CHECK_FILL = PatternFill("solid", fgColor="FFF2CC")


def cell(v):
    """Значение ячейки: скаляры как есть, списки скаляров — через «; », остальное — JSON одной строкой."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        return v
    if isinstance(v, list) and all(isinstance(x, (str, int, float)) for x in v):
        return "; ".join(str(x) for x in v)
    return json.dumps(v, ensure_ascii=False)


def collections(data, name):
    """Таблицы файла: список словарей верхнего уровня или ключи-списки/словари словарей. → [(имя, [строки-словари])]."""
    out = []
    if isinstance(data, list):
        if data and all(isinstance(x, dict) for x in data):
            out.append((name, data))
        elif data:
            out.append((name, [{"значение": x} for x in data]))
        return out
    if not isinstance(data, dict):
        return out
    for k, v in data.items():
        if k.startswith("_"):
            continue
        if isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
            out.append((f"{name}·{k}", v))
        elif isinstance(v, list) and v and all(isinstance(x, list) for x in v):
            width = max(len(x) for x in v)
            out.append(
                (
                    f"{name}·{k}",
                    [
                        {
                            f"колонка {i + 1}": x[i] if i < len(x) else ""
                            for i in range(width)
                        }
                        for x in v
                    ],
                )
            )
        elif isinstance(v, dict) and v and all(isinstance(x, dict) for x in v.values()):
            out.append((f"{name}·{k}", [{"ключ": kk, **vv} for kk, vv in v.items()]))
        elif isinstance(v, dict) and v:
            out.append(
                (f"{name}·{k}", [{"ключ": kk, "значение": vv} for kk, vv in v.items()])
            )
    if not out:  # плоский словарь верхнего уровня
        rows = [
            {"ключ": k, "значение": v} for k, v in data.items() if not k.startswith("_")
        ]
        if rows:
            out.append((name, rows))
    return out


def sheet_name(raw, used):
    s = re.sub(r"[\[\]\*\?/\\:]", "-", raw)[:31] or "лист"
    base, i = s, 2
    while s in used:
        s = f"{base[:28]}-{i}"
        i += 1
    used.add(s)
    return s


def write_sheet(ws, rows):
    cols = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    header = cols + [CHECK_COL]
    ws.append(header)
    for r in rows:
        ws.append([cell(r.get(c)) for c in cols] + [""])
    for i, name in enumerate(header, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = CHECK_FILL if name == CHECK_COL else HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        width = min(
            60,
            max(
                10,
                len(name) + 2,
                *(
                    len(str(ws.cell(row=j, column=i).value or "")) + 2
                    for j in range(2, min(ws.max_row, 60) + 1)
                ),
            ),
        )
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def passport_rows():
    rows = []
    for d, status in PASSPORT_DIRS:
        for f in sorted((SEED / d).glob("M-*.json")):
            p = json.loads(f.read_text("utf-8"))
            src = p.get("sources") or {}
            value = p.get("value") or {}
            ext = p.get("extractor") or {}
            rows.append(
                {
                    "код": p.get("code", f.stem),
                    "статус": status,
                    "название": p.get("title"),
                    "версия": p.get("version"),
                    "суть проверки": p.get("summary"),
                    "основание (Матрица)": p.get("basis"),
                    "вид значения": value.get("kind"),
                    "единица": value.get("unit"),
                    "извлечение": ext.get("kind"),
                    "источники ПД": "; ".join(
                        s.get("discipline", "")
                        for s in src.get("PD", [])
                        if isinstance(s, dict)
                    ),
                    "источники РД": "; ".join(
                        s.get("discipline", "")
                        for s in src.get("RD", [])
                        if isinstance(s, dict)
                    ),
                    "источники ИД": "; ".join(
                        s.get("discipline", "")
                        for s in src.get("ID", [])
                        if isinstance(s, dict)
                    ),
                    "файл": str(f.relative_to(ROOT)),
                }
            )
    return rows


def main() -> int:
    wb = Workbook()
    toc = wb.active
    toc.title = "Оглавление"
    used = {"Оглавление"}
    toc_rows = []

    rows = passport_rows()
    ws = wb.create_sheet(sheet_name("Паспорта параметров", used))
    write_sheet(ws, rows)
    toc_rows.append(
        {
            "лист": ws.title,
            "источник": "data/seed/passports*/M-*.json",
            "что это": "Паспорта параметров: единый источник для извлечения, выбора источника и сравнения",
            "строк": len(rows),
        }
    )

    for fname, about in FILES:
        path = SEED / fname
        if not path.exists():
            print(f"нет файла: {fname}", file=sys.stderr)
            continue
        data = json.loads(path.read_text("utf-8"))
        tables = collections(data, path.stem)
        note = data.get("_about", "") if isinstance(data, dict) else ""
        for tname, trows in tables:
            ws = wb.create_sheet(sheet_name(tname, used))
            write_sheet(ws, trows)
            toc_rows.append(
                {
                    "лист": ws.title,
                    "источник": f"data/seed/{fname}",
                    "что это": about,
                    "строк": len(trows),
                    "заметка файла": note,
                }
            )

    write_sheet(toc, toc_rows)
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / "справочники.xlsx"
    wb.save(target)
    print(
        f"{target.relative_to(ROOT)}: {len(wb.sheetnames)} листов, {sum(r['строк'] for r in toc_rows)} строк"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
